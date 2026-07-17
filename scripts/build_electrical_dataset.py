"""Build an audited electrical dataset from the original instrument exports.

The instrument tables contain wall-clock timestamps, voltage, a current field,
and active power.  Some exports label the current field as mA even though the
numeric values satisfy P ~= V * I only when interpreted as amperes.  This
builder therefore infers the scale from voltage-current-power consistency and
records any header conflict instead of silently trusting the unit label.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # XLSX cross-checking is optional at import time.
    load_workbook = None


DATE_RE = re.compile(r"(20\d{2})-(\d{2})-(\d{2})")
UNIT_RE = re.compile(r"\(([^()]*)\)")


def _split_markdown_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def _normalise_header(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().lower())


def _find_column(header: list[object], aliases: tuple[str, ...], required: bool = True) -> int | None:
    normalised = [_normalise_header(value) for value in header]
    for index, token in enumerate(normalised):
        if any(alias in token for alias in aliases):
            return index
    if required:
        raise ValueError(f"Required column not found; aliases={aliases}, header={header}")
    return None


def _column_map(header: list[object]) -> dict[str, int | None]:
    return {
        "sequence": _find_column(header, ("序号", "sequence", "index")),
        "product_code": _find_column(header, ("产品编号", "product"), required=False),
        "instrument_env_C": _find_column(header, ("环境温度", "ambienttemperature"), required=False),
        "timestamp": _find_column(header, ("测试时间", "timestamp", "testtime")),
        "voltage_V": _find_column(header, ("电压", "voltage")),
        "current_raw": _find_column(header, ("电流", "current")),
        "active_power_W": _find_column(header, ("有功功率", "activepower")),
        "power_factor_raw": _find_column(header, ("功率因数", "powerfactor"), required=False),
        "frequency_Hz": _find_column(header, ("频率", "frequency"), required=False),
        "apparent_power_W": _find_column(header, ("视在功率", "apparentpower"), required=False),
        "reactive_power_var": _find_column(header, ("无功功率", "reactivepower"), required=False),
    }


def _value(row: list[object], index: int | None, default: object = None) -> object:
    if index is None or index >= len(row):
        return default
    value = row[index]
    return default if value in (None, "") else value


def _float_or_nan(value: object) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return math.nan


def _records_from_rows(rows: Iterable[list[object]], source_file: str) -> tuple[pd.DataFrame, str]:
    materialised = list(rows)
    header_index = None
    header = None
    for index, row in enumerate(materialised[:20]):
        tokens = [_normalise_header(value) for value in row]
        if any("测试时间" in token or "timestamp" in token for token in tokens) and any(
            "有功功率" in token or "activepower" in token for token in tokens
        ):
            header_index = index
            header = row
            break
    if header is None or header_index is None:
        raise ValueError(f"Instrument header not found in {source_file}")

    columns = _column_map(header)
    current_header = str(header[columns["current_raw"]])
    unit_match = UNIT_RE.search(current_header)
    current_header_unit = unit_match.group(1) if unit_match else "unknown"

    records = []
    for row in materialised[header_index + 1 :]:
        sequence_value = _value(row, columns["sequence"])
        try:
            sequence = int(float(sequence_value))
        except (TypeError, ValueError):
            continue
        timestamp = pd.to_datetime(_value(row, columns["timestamp"]), errors="coerce")
        if pd.isna(timestamp):
            continue
        records.append(
            {
                "sequence": sequence,
                "product_code": str(_value(row, columns["product_code"], "")),
                "instrument_env_C": _float_or_nan(_value(row, columns["instrument_env_C"])),
                "timestamp": timestamp,
                "voltage_V": _float_or_nan(_value(row, columns["voltage_V"])),
                "current_raw": _float_or_nan(_value(row, columns["current_raw"])),
                "active_power_W": _float_or_nan(_value(row, columns["active_power_W"])),
                "power_factor_raw": _float_or_nan(_value(row, columns["power_factor_raw"])),
                "frequency_Hz": _float_or_nan(_value(row, columns["frequency_Hz"])),
                "apparent_power_W": _float_or_nan(_value(row, columns["apparent_power_W"])),
                "reactive_power_var": _float_or_nan(_value(row, columns["reactive_power_var"])),
                "source_file": source_file,
            }
        )
    frame = pd.DataFrame.from_records(records)
    if frame.empty:
        raise ValueError(f"No instrument records parsed from {source_file}")
    return frame, current_header_unit


def load_power_markdown(path: Path) -> tuple[pd.DataFrame, str]:
    rows = [_split_markdown_row(line) for line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines()]
    return _records_from_rows(rows, path.name)


def load_power_xlsx(path: Path) -> tuple[pd.DataFrame, str]:
    if load_workbook is None:
        raise ModuleNotFoundError("openpyxl is required for XLSX cross-checking")
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    return _records_from_rows(rows, path.name)


def _infer_current_scale(frame: pd.DataFrame) -> tuple[float, dict[str, float]]:
    usable = frame[
        frame["voltage_V"].gt(0.1)
        & frame["current_raw"].gt(1e-6)
        & frame["active_power_W"].gt(0.1)
    ]
    if len(usable) < 3:
        raise ValueError("Too few active rows to infer the electrical-current scale")

    candidate_errors = {}
    for scale in (1.0, 0.001):
        predicted = usable["voltage_V"] * usable["current_raw"] * scale
        relative = (predicted - usable["active_power_W"]).abs() / usable["active_power_W"].clip(lower=1.0)
        candidate_errors[str(scale)] = float(relative.median())
    selected = min((1.0, 0.001), key=lambda scale: candidate_errors[str(scale)])
    if candidate_errors[str(selected)] > 0.10:
        raise ValueError(f"No plausible current scale satisfies P ~= V*I: {candidate_errors}")
    return selected, candidate_errors


def _prepare_day(frame: pd.DataFrame, current_header_unit: str, source_sha256: str) -> tuple[pd.DataFrame, dict]:
    frame = frame.copy().reset_index(drop=True)
    frame.insert(0, "source_row_order", np.arange(len(frame), dtype=int))
    raw_time_deltas = frame["timestamp"].diff().dt.total_seconds().dropna()
    raw_sequence_deltas = frame["sequence"].diff().dropna()
    raw_sequence_first_last = [int(frame["sequence"].iloc[0]), int(frame["sequence"].iloc[-1])]
    frame = frame.sort_values(["timestamp", "sequence"], kind="stable").reset_index(drop=True)
    scale, scale_errors = _infer_current_scale(frame)
    frame.insert(0, "experiment_date", frame["timestamp"].dt.strftime("%Y-%m-%d"))
    frame.insert(1, "sample_order", np.arange(len(frame), dtype=int))
    frame["elapsed_s"] = (frame["timestamp"] - frame["timestamp"].iloc[0]).dt.total_seconds()
    frame["current_header_unit"] = current_header_unit
    frame["current_scale_to_A"] = scale
    frame["current_A"] = frame["current_raw"] * scale
    frame["calculated_power_W"] = frame["voltage_V"] * frame["current_A"]
    denominator = frame["active_power_W"].abs().clip(lower=1.0)
    frame["vi_relative_error"] = (frame["calculated_power_W"] - frame["active_power_W"]).abs() / denominator

    near_zero_measured = frame["active_power_W"].abs().le(0.1)
    near_zero_calculated = frame["calculated_power_W"].abs().le(0.1)
    frame["vi_consistent"] = (near_zero_measured & near_zero_calculated) | (
        ~near_zero_measured & frame["vi_relative_error"].le(0.02)
    )
    finite = np.isfinite(frame[["voltage_V", "current_A", "active_power_W"]]).all(axis=1)
    plausible = (
        frame["voltage_V"].between(0.0, 60.0)
        & frame["current_A"].between(0.0, 10.0)
        & frame["active_power_W"].between(0.0, 150.0)
    )
    frame["valid_electrical_range"] = finite & plausible
    frame["duplicate_timestamp"] = frame["timestamp"].duplicated(keep=False)
    frame["valid_for_model"] = frame["valid_electrical_range"] & frame["vi_consistent"]
    frame["source_sha256"] = source_sha256

    header_normalised = current_header_unit.strip().lower().replace(" ", "")
    header_implied_scale = 0.001 if header_normalised == "ma" else 1.0 if header_normalised == "a" else None
    header_conflict = header_implied_scale is not None and not math.isclose(header_implied_scale, scale)
    deltas = frame["timestamp"].diff().dt.total_seconds().dropna()
    median_dt = float(deltas[deltas > 0].median())
    active_error = frame.loc[frame["active_power_W"].gt(1.0), "vi_relative_error"]
    qc = {
        "source_file": str(frame["source_file"].iloc[0]),
        "source_sha256": source_sha256,
        "experiment_date": str(frame["experiment_date"].iloc[0]),
        "rows": int(len(frame)),
        "sequence_first_last": raw_sequence_first_last,
        "sequence_gaps": int((raw_sequence_deltas != 1).sum()),
        "raw_order_negative_time_steps": int((raw_time_deltas < 0).sum()),
        "raw_order_duplicate_time_steps": int((raw_time_deltas == 0).sum()),
        "raw_order_is_monotonic": bool((raw_time_deltas >= 0).all()),
        "timestamp_start": frame["timestamp"].iloc[0].isoformat(),
        "timestamp_end": frame["timestamp"].iloc[-1].isoformat(),
        "duration_h": round(float(frame["elapsed_s"].iloc[-1]) / 3600.0, 4),
        "median_sampling_s": round(median_dt, 4),
        "duplicate_timestamp_extra_rows": int(frame["timestamp"].duplicated().sum()),
        "nonpositive_time_steps": int((deltas <= 0).sum()),
        "gaps_over_2x_median": int((deltas > 2.0 * median_dt).sum()),
        "current_header_unit": current_header_unit,
        "current_scale_to_A": scale,
        "current_scale_candidate_median_errors": scale_errors,
        "current_header_conflict": bool(header_conflict),
        "voltage_range_V": [float(frame["voltage_V"].min()), float(frame["voltage_V"].max())],
        "current_range_A": [float(frame["current_A"].min()), float(frame["current_A"].max())],
        "active_power_range_W": [float(frame["active_power_W"].min()), float(frame["active_power_W"].max())],
        "zero_power_rows": int(frame["active_power_W"].eq(0.0).sum()),
        "invalid_electrical_range_rows": int((~frame["valid_electrical_range"]).sum()),
        "vi_inconsistent_rows": int((~frame["vi_consistent"]).sum()),
        "vi_relative_error_median": float(active_error.median()),
        "vi_relative_error_p95": float(active_error.quantile(0.95)),
        "valid_for_model_rows": int(frame["valid_for_model"].sum()),
    }
    return frame, qc


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_electrical_dataset(power_dir: Path) -> tuple[pd.DataFrame, dict]:
    frames = []
    day_qc = {}
    files = sorted(power_dir.glob("*.md"))
    for path in files:
        try:
            raw, current_header_unit = load_power_markdown(path)
        except ValueError:
            continue
        prepared, qc = _prepare_day(raw, current_header_unit, _sha256(path))
        date = qc["experiment_date"]
        if date in day_qc:
            raise ValueError(f"More than one electrical Markdown file parsed for {date}")
        frames.append(prepared)
        day_qc[date] = qc
    if not frames:
        raise ValueError(f"No electrical instrument tables found in {power_dir}")

    dataset = pd.concat(frames, ignore_index=True)
    qc = {
        "source": "original electrical-parameter instrument Markdown exports",
        "rows_total": int(len(dataset)),
        "rows_valid_for_model": int(dataset["valid_for_model"].sum()),
        "days_parsed": len(day_qc),
        "days": day_qc,
        "known_limitations": [
            "The current header is inconsistent with the numeric scale on some days; scale is inferred from P ~= V*I and the conflict is retained in QC.",
            "Duplicate wall-clock timestamps are preserved and flagged rather than silently removed.",
            "The instrument's constant environment-temperature field is metadata, not a measured outdoor ambient-temperature series.",
            "Electrical timestamps cannot be aligned to relative-time thermal logs without an explicit or auditable offset procedure.",
        ],
    }
    return dataset, qc


def crosscheck_xlsx(dataset: pd.DataFrame, xlsx_dir: Path) -> dict:
    checks = {}
    for path in sorted(xlsx_dir.glob("*.xlsx")):
        raw, header_unit = load_power_xlsx(path)
        prepared, _ = _prepare_day(raw, header_unit, _sha256(path))
        date = str(prepared["experiment_date"].iloc[0])
        markdown = dataset[dataset["experiment_date"] == date].reset_index(drop=True)
        if markdown.empty:
            checks[date] = {"xlsx_file": path.name, "matched_markdown": False}
            continue
        numeric_columns = ["voltage_V", "current_raw", "active_power_W"]
        same_rows = len(markdown) == len(prepared)
        checks[date] = {
            "xlsx_file": path.name,
            "matched_markdown": True,
            "rows_xlsx": int(len(prepared)),
            "rows_markdown": int(len(markdown)),
            "row_count_equal": bool(same_rows),
            "timestamps_equal": bool(same_rows and markdown["timestamp"].equals(prepared["timestamp"])),
            "max_absolute_numeric_difference": {
                column: float((markdown[column] - prepared[column]).abs().max()) if same_rows else None
                for column in numeric_columns
            },
        }
    return checks


def _write_qc_markdown(qc: dict, path: Path) -> None:
    lines = [
        "# Electrical experiment data quality report",
        "",
        f"Parsed {qc['rows_total']:,} rows across {qc['days_parsed']} experiment days; "
        f"{qc['rows_valid_for_model']:,} rows pass the current range and V-I-P consistency flags.",
        "",
        "| Date | Rows | Start | End | Median dt (s) | P range (W) | Duplicate timestamps | V-I-P inconsistent | Header conflict |",
        "|---|---:|---|---|---:|---:|---:|---:|:---:|",
    ]
    for date, values in qc["days"].items():
        power_range = values["active_power_range_W"]
        lines.append(
            f"| {date} | {values['rows']} | {values['timestamp_start'][11:19]} | {values['timestamp_end'][11:19]} | "
            f"{values['median_sampling_s']:.1f} | {power_range[0]:.2f} to {power_range[1]:.2f} | "
            f"{values['duplicate_timestamp_extra_rows']} | {values['vi_inconsistent_rows']} | "
            f"{'yes' if values['current_header_conflict'] else 'no'} |"
        )
    if qc.get("xlsx_crosscheck"):
        lines.extend(["", "## XLSX cross-check", ""])
        for date, check in qc["xlsx_crosscheck"].items():
            lines.append(
                f"- {date}: row count equal={check.get('row_count_equal')}; "
                f"timestamps equal={check.get('timestamps_equal')}; "
                f"max numeric differences={check.get('max_absolute_numeric_difference')}."
            )
    lines.extend(["", "## Known limitations", ""])
    lines.extend(f"- {item}" for item in qc["known_limitations"])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--power-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--xlsx-dir", type=Path)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    dataset, qc = build_electrical_dataset(args.power_dir)
    if args.xlsx_dir:
        qc["xlsx_crosscheck"] = crosscheck_xlsx(dataset, args.xlsx_dir)
    dataset.to_csv(args.output_dir / "electrical_clean_v1.csv", index=False)
    (args.output_dir / "electrical_clean_v1_qc.json").write_text(
        json.dumps(qc, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    _write_qc_markdown(qc, args.output_dir / "electrical_clean_v1_qc.md")
    print(
        json.dumps(
            {
                "rows": len(dataset),
                "valid_for_model": int(dataset["valid_for_model"].sum()),
                "days": qc["days_parsed"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
